"""Regenerate the two maintained Italian Word guides from their Markdown sources.

Run with Python containing python-docx (the Codex document runtime supplies it).
Document tooling is deliberately separate from application dependencies.
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
GUIDES = (
    ("cliente/GUIDA_CLIENTE", "TRD · Guida operativa cliente"),
    ("team/GUIDA_PREPARAZIONE", "TRD · Preparazione del team"),
)


def inline(paragraph, value):
    for part in re.split(r"(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)", value):
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            target = paragraph.part.relate_to(
                link[2],
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), target)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            style = OxmlElement("w:rStyle")
            style.set(qn("w:val"), "Hyperlink")
            properties.append(style)
            run.append(properties)
            text = OxmlElement("w:t")
            text.text = link[1]
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Liberation Mono"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def table(document, rows):
    count = len(rows[0])
    result = document.add_table(rows=0, cols=count)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    widths = [Cm(7.6), Cm(9.2)] if count == 2 else [Cm(4.9), Cm(6.5), Cm(5.4)]
    for column, width in zip(result.columns, widths, strict=True):
        column.width = width
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), val)
        borders.append(border)
    result._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = result.add_row()
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if index == 0:
            properties.append(OxmlElement("w:tblHeader"))
        for cell, width, value in zip(row.cells, widths, values, strict=True):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                margin = OxmlElement(f"w:{edge}")
                margin.set(qn("w:w"), "95")
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.style = "Guide Table"
            inline(paragraph, value)
            for run in paragraph.runs:
                run.bold = index == 0
            if index == 0 or index % 2 == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "DFEAF3" if index == 0 else "F5F8FA")
                cell._tc.get_or_add_tcPr().append(shade)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build(stem, label):
    source = ROOT / "docs" / f"{stem}.md"
    destination = source.with_suffix(".docx")
    lines = source.read_text(encoding="utf-8").splitlines()
    updated = re.search(r"Aggiornamento del (.*?) \|", "\n".join(lines))[1]
    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(2)
    section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
    section.header_distance, section.footer_distance = Cm(0.8), Cm(0.8)
    styles = document.styles
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        styles[name].font.name = "Calibri"
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    normal = styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.keep_together = True
    for border in styles["Title"].element.findall(".//" + qn("w:pBdr")):
        border.getparent().remove(border)
    styles["Title"].font.size = Pt(24)
    styles["Title"].paragraph_format.space_after = Pt(10)
    styles["Subtitle"].font.size = Pt(10)
    styles["Subtitle"].paragraph_format.space_after = Pt(14)
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].paragraph_format.space_before = Pt(12)
        styles[name].paragraph_format.space_after = Pt(6)
        styles[name].paragraph_format.keep_with_next = True
    table_style = styles.add_style("Guide Table", 1)
    table_style.base_style = normal
    table_style.font.size = Pt(9.4)
    table_style.paragraph_format.space_after = Pt(2)
    table_style.paragraph_format.line_spacing = 1.04
    code_style = styles.add_style("Guide Code", 1)
    code_style.base_style = normal
    code_style.font.name = "Liberation Mono"
    code_style.font.size = Pt(8.5)
    code_style.paragraph_format.space_after = Pt(3)
    revision_style = styles.add_style("Guide Revision", 1)
    revision_style.base_style = normal
    revision_style.paragraph_format.line_spacing = 1
    revision_style.paragraph_format.space_after = Pt(2)
    section.header.paragraphs[0].text = label
    section.header.paragraphs[0].runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.add_run(f"Aggiornamento {updated} · Pagina ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = label
    document.core_properties.author = "TRD"
    document.core_properties.language = "it-IT"
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "it-IT")
    normal.element.get_or_add_rPr().append(language)
    index = 0
    first_paragraph = True
    in_revisions = False
    list_number_id = None
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if not line:
            continue
        if line.startswith("# "):
            document.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            in_revisions = line == "## Cronologia delle revisioni"
            document.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("### "):
            document.add_paragraph(line[4:], "Heading 2")
        elif image := re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line):
            image_path = (source.parent / image[2]).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            picture = paragraph.add_run().add_picture(str(image_path), width=Cm(16.8))
            picture._inline.docPr.set("descr", image[1])
        elif line.startswith("```"):
            while index < len(lines) and not lines[index].startswith("```"):
                document.add_paragraph(lines[index], "Guide Code")
                index += 1
            index += 1
        elif line.startswith("|"):
            rows = [[value.strip() for value in line.strip("|").split("|")]]
            while index < len(lines) and lines[index].startswith("|"):
                values = [value.strip() for value in lines[index].strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", value) for value in values):
                    rows.append(values)
                index += 1
            table(document, rows)
        elif re.match(r"\d+\. ", line):
            paragraph = document.add_paragraph(style="List Number")
            if line.startswith("1. ") or list_number_id is None:
                numbering = document.part.numbering_part.element
                base_id = styles["List Number"].element.pPr.numPr.numId.val
                abstract_id = numbering.xpath(f'./w:num[@w:numId="{base_id}"]/w:abstractNumId')[
                    0
                ].get(qn("w:val"))
                new_number = numbering.add_num(int(abstract_id))
                new_number.add_lvlOverride(ilvl=0).add_startOverride(1)
                list_number_id = new_number.numId
            paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = list_number_id
            paragraph.paragraph_format.keep_with_next = bool(
                index < len(lines) and re.match(r"\d+\. ", lines[index])
            )
            inline(paragraph, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            inline(paragraph, line[2:])
        else:
            text = line
            while (
                index < len(lines)
                and lines[index].strip()
                and not lines[index].startswith(("#", "|", "```", "![", "- "))
                and not re.match(r"\d+\. ", lines[index])
            ):
                text += " " + lines[index].strip()
                index += 1
            style = "Subtitle" if first_paragraph else "Guide Revision" if in_revisions else None
            inline(document.add_paragraph(style=style), text)
            first_paragraph = False
    document.save(destination)
    print(destination.relative_to(ROOT))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--guide", choices=("cliente", "team"), help="Build only the affected guide"
    )
    args = parser.parse_args()
    for guide in GUIDES:
        if args.guide is None or guide[0].startswith(args.guide + "/"):
            build(*guide)
